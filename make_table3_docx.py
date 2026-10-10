#!/usr/bin/env python3
"""
Generate Table 3 as a formatted Word document from the CSV output
of lca_table3.py.

Usage:
    python make_table3_docx.py [--input table3.csv] [--output Table_3.docx]

Dependencies: pandas, python-docx
"""

import argparse
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cell_font(
    cell, text, bold=False, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, italic=False
):
    """Set cell text with formatting."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    run.bold = bold
    run.italic = italic
    # Vertical alignment
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = qn("w:vAlign")
    va = tcPr.find(vAlign)
    if va is None:
        va = tcPr.makeelement(vAlign, {})
        tcPr.append(va)
    va.set(qn("w:val"), "center")


def main():
    parser = argparse.ArgumentParser(
        description="Generate formatted Table 3 Word document"
    )
    parser.add_argument(
        "--input",
        "-i",
        default="table3_baseline_by_class.csv",
        help="Input CSV from lca_table3.py",
    )
    parser.add_argument(
        "--output", "-o", default="Table_3_baseline_by_class.docx", help="Output .docx"
    )
    args = parser.parse_args()

    df = pd.read_csv(args.input, index_col=0)
    col_order = [
        "Well-Informed",
        "Selectively Skeptical",
        "Uncertain/Disengaged",
        "Misinformed",
    ]
    df = df[col_order + ["P value"]]

    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Table 3. ")
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Times New Roman"
    run = title.add_run(
        "Baseline (Wave 27) characteristics by latent class, estimated using "
        "posterior-weighted proportions and means from the four-class common-item "
        "latent class model."
    )
    run.font.size = Pt(11)
    run.font.name = "Times New Roman"

    # Row definitions
    rows_data = [
        ("header", "Demographics"),
        ("data", "N (% of sample)", "N"),
        ("data", "Male (%)", "Male (%)"),
        ("data", "Female (%)", "Female (%)"),
        ("data", "Age 18-34 (%)", "Age 18-34 (%)"),
        ("data", "Age 35-54 (%)", "Age 35-54 (%)"),
        ("data", "Age 55-64 (%)", "Age 55-64 (%)"),
        ("data", "Age ≥ 65 (%)", "Age >= 65 (%)"),
        (
            "data",
            "Bachelor’s or postgraduate education (%)",
            "Bachelor's or postgraduate (%)",
        ),
        ("data", "Rural community (%)", "Rural community (%)"),
        (
            "data",
            "Poor or working social class (%)",
            "Poor or working social class (%)",
        ),
        ("header", "Vaccination status"),
        ("data", "Unvaccinated (%)", "Unvaccinated (%)"),
        ("data", "Three or more doses (%)", "Three or more vaccine doses (%)"),
        ("header", "Trust (1-7 scale), mean (SD)"),
        ("data", "Doctors", "Trust in doctors (1-7)"),
        ("data", "Scientists", "Trust in scientists (1-7)"),
        ("data", "Public health experts", "Trust in public health experts (1-7)"),
        ("data", "Journalists", "Trust in journalists (1-7)"),
        ("header", "Information sources and attitudes"),
        (
            "data",
            "Approval of PHAC (1-5), mean (SD)",
            "Approval of PHAC (1-5)",
        ),
        (
            "data",
            "Reliance on health professionals (1-5), mean (SD)",
            "Reliance on health professionals (1-5)",
        ),
        (
            "data",
            "Reliance on public health agencies (1-5), mean (SD)",
            "Reliance on public health agencies (1-5)",
        ),
        (
            "data",
            "Reliance on traditional media (1-5), mean (SD)",
            "Reliance on traditional media (1-5)",
        ),
        (
            "data",
            "Reliance on social media (1-5), mean (SD)",
            "Reliance on social media (1-5)",
        ),
        (
            "data",
            "Concern about disinformation (1-5), mean (SD)",
            "Concern about disinformation (1-5)",
        ),
    ]

    n_cols = 6
    table = doc.add_table(rows=0, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    col_widths = [
        Inches(2.4),
        Inches(0.9),
        Inches(0.9),
        Inches(0.9),
        Inches(0.9),
        Inches(0.7),
    ]

    # Header row
    hdr = table.add_row()
    set_cell_font(
        hdr.cells[0], "Characteristic", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT
    )
    for i, cls in enumerate(col_order):
        set_cell_font(hdr.cells[i + 1], cls, bold=True)
    set_cell_font(hdr.cells[5], "P value", bold=True, italic=True)
    for i, w in enumerate(col_widths):
        hdr.cells[i].width = w

    # Data rows
    for row_spec in rows_data:
        row = table.add_row()
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

        if row_spec[0] == "header":
            set_cell_font(
                row.cells[0], row_spec[1], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT
            )
            for i in range(1, n_cols):
                set_cell_font(row.cells[i], "")
            # Light gray shading
            for i in range(n_cols):
                tc = row.cells[i]._tc
                tcPr = tc.get_or_add_tcPr()
                shading = tcPr.makeelement(
                    qn("w:shd"),
                    {
                        qn("w:val"): "clear",
                        qn("w:color"): "auto",
                        qn("w:fill"): "F2F2F2",
                    },
                )
                tcPr.append(shading)
        else:
            label = row_spec[1]
            csv_key = row_spec[2]
            set_cell_font(
                row.cells[0], label, align=WD_ALIGN_PARAGRAPH.LEFT
            )
            for i, cls in enumerate(col_order):
                val = df.loc[csv_key, cls] if csv_key in df.index else ""
                set_cell_font(row.cells[i + 1], val)
            pval = df.loc[csv_key, "P value"] if csv_key in df.index else ""
            if pd.isna(pval):
                pval = ""
            set_cell_font(row.cells[5], str(pval))

    # Footnote
    footnote = doc.add_paragraph()
    footnote.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = footnote.add_run(
        "Note: Estimates are posterior-weighted from the four-class latent class model "
        "fitted to the four common misinformation items (COVID-19 deaths exaggerated, "
        "vaccine deaths hidden, vaccines cause infertility, vaccines alter DNA). "
        "N = 1,185 respondents with complete data on all four items. "
        "P values are from posterior-weighted Wald tests: chi-square tests for "
        "categorical variables and F-tests for continuous variables. "
        "Trust scales range from 1 (no trust at all) to 7 (a great deal of trust). "
        "Reliance and approval scales range from 1 to 5. "
        "PHAC = Public Health Agency of Canada."
    )
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"
    run.italic = True

    doc.save(args.output)
    print(f"Saved to {args.output}")


if __name__ == "__main__":
    main()
